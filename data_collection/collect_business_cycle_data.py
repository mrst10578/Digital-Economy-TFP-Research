# -*- coding: utf-8 -*-
"""
Phase 1 data collector for:
Digital Economy, Artificial Intelligence, and Business Cycle Dynamics

Revision 2 (2026-09-28)
- keeps Iran in cross-country auditing;
- removes the artificial World Bank 2000 start-year floor;
- paginates World Bank responses and labels countries vs aggregates;
- adds missing WDI digital indicators from the research document;
- uses an official FRED CSV fallback when no FRED API key is configured;
- chunks BLS history instead of discarding older observations;
- uses IMF AI Preparedness indicator AI_PI rather than dataset name AIPI;
- uses BEA's latest Digital Economy 2017-2022 workbook;
- writes coverage, Iran data, source dictionary, and collection status to Excel.

This script collects and audits data only. It does not run econometric models.
"""

from __future__ import annotations

import csv
import io
import os
import re
import time
import zipfile
from datetime import datetime
from urllib.parse import urljoin

import pandas as pd
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

OUTPUT_FILE = "AI_Digital_Economy_BusinessCycle_Data.xlsx"
TEMP_OUTPUT_FILE = OUTPUT_FILE + ".tmp.xlsx"
END_YEAR = datetime.now().year
REQUEST_TIMEOUT = 60
SLEEP_BETWEEN_CALLS = 0.25
RETRIEVED_AT = datetime.utcnow().replace(microsecond=0).isoformat() + "Z"

FRED_API_KEY = os.environ.get("FRED_API_KEY", "")
BEA_API_KEY = os.environ.get("BEA_API_KEY", "")
BLS_API_KEY = os.environ.get("BLS_API_KEY", "")

_log_rows: list[dict] = []
_coverage_rows: list[dict] = []
_iran_rows: list[dict] = []
_dictionary_rows: list[dict] = []
_country_name_to_iso3: dict[str, str] = {}

retry = Retry(
    total=3,
    connect=3,
    read=3,
    backoff_factor=1.0,
    status_forcelist=(429, 500, 502, 503, 504),
    allowed_methods=frozenset({"GET", "POST"}),
)
SESSION = requests.Session()
SESSION.mount("https://", HTTPAdapter(max_retries=retry))


def log(name: str, source: str, status: str, detail: str = "", rows: int | None = None) -> None:
    _log_rows.append({
        "Variable": name,
        "Source": source,
        "Status": status,
        "Rows": rows,
        "Detail": clean_excel_text(str(detail)) if detail is not None else "",
        "Retrieved_At": RETRIEVED_AT,
    })
    suffix = f" ({detail})" if detail else ""
    print(f"[{status}] {source} :: {name}{suffix}")


def register(
    variable_name: str,
    source: str,
    series_id: str,
    geography: str,
    frequency: str = "",
    requirement_role: str = "",
    note: str = "",
) -> None:
    _dictionary_rows.append({
        "Variable": variable_name,
        "Source": source,
        "Series_ID": series_id,
        "Geography": geography,
        "Frequency_Hint": frequency,
        "Requirement_Role": requirement_role,
        "Note": note,
    })


def safe_sheet_name(name: str) -> str:
    for ch in '[]:*?/\\':
        name = name.replace(ch, "")
    return name[:31]


_ILLEGAL_EXCEL_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def clean_excel_text(value):
    if isinstance(value, str):
        return _ILLEGAL_EXCEL_CHARS.sub("", value)[:32767]
    return value


def excel_safe_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_object_dtype(out[col].dtype) or pd.api.types.is_string_dtype(out[col].dtype):
            out[col] = out[col].map(clean_excel_text)
    return out


def write_sheet(writer, df: pd.DataFrame | None, name: str, source: str) -> None:
    if df is None or df.empty:
        log(name, source, "EMPTY")
        return
    excel_safe_dataframe(df).to_excel(writer, sheet_name=safe_sheet_name(name), index=False)
    log(name, source, "OK", f"{len(df)} rows", len(df))


def numeric_valid(series: pd.Series) -> pd.Series:
    return pd.to_numeric(series, errors="coerce").notna()


def add_iran_long_rows(
    df: pd.DataFrame,
    source: str,
    series_id: str,
    period_col: str,
    value_col: str,
    country_col: str,
    observation_status: str,
) -> None:
    if country_col not in df.columns or value_col not in df.columns:
        return
    subset = df[df[country_col].astype(str) == "IRN"].copy()
    subset = subset[numeric_valid(subset[value_col])].copy()
    for _, row in subset.iterrows():
        _iran_rows.append({
            "source": source,
            "series_id": series_id,
            "country_iso3": "IRN",
            "period": row.get(period_col),
            "value": row.get(value_col),
            "observation_status": observation_status,
            "retrieved_at": RETRIEVED_AT,
        })


# ---------------------------------------------------------------------------
# FRED
# ---------------------------------------------------------------------------

FRED_SERIES = {
    "Real_GDP_Level": "GDPC1",
    "Real_GDP_Growth_Rate": "A191RL1Q225SBEA",
    "Industrial_Production": "INDPRO",
    "NBER_Recession_Indicator": "USREC",
    "Sahm_Rule_Indicator": "SAHMREALTIME",
    "Yield_Curve_10Y_minus_2Y": "T10Y2Y",
    "Nonfarm_Payrolls": "PAYEMS",
    "Unemployment_Rate": "UNRATE",
    "VIX_Volatility_Index": "VIXCLS",
    "OECD_CLI_United_States": "USALOLITONOSTSAM",
    "Info_Processing_Equip_Software_Inv": "A679RC1Q027SBEA",
    "Info_Processing_Inv_Contribution": "A679RZ2Q224SBEA",
    "Computer_Systems_Design_Employment": "CES6054150001",
    "Information_Sector_Employment": "USINFO",
    "JOLTS_Job_Openings_Total": "JTSJOL",
    "JOLTS_Prof_Business_Services": "JTS540099JOL",
    "JOLTS_Hires_Total": "JTSHIL",
    "JOLTS_Quits_Total": "JTSQUL",
    "JOLTS_Total_Separations": "JTSTSL",
    "Labor_Productivity_NonfarmBiz": "OPHNFB",
    "Unit_Labor_Cost_NonfarmBiz": "ULCNFB",
    "Total_Factor_Productivity": "MFPPBS",
    "CPI_Headline": "CPIAUCSL",
    "CPI_Core": "CPILFESL",
    "PCE_Price_Index": "PCEPI",
    "PCE_Core": "PCEPILFE",
    "Economic_Policy_Uncertainty_Monthly": "USEPUINDXM",
    "Economic_Policy_Uncertainty_Daily": "USEPUINDXD",
    "National_Financial_Conditions_Index": "NFCI",
    "Effective_Federal_Funds_Rate_Candidate": "FEDFUNDS",
    "SP500_Equity_Index": "SP500",
    "US_10Y_Treasury_Yield": "DGS10",
    "WTI_Crude_Oil_Price": "DCOILWTICO",
    "Broad_US_Dollar_Index": "DTWEXBGS",
    "Nonresidential_Structures_Investment": "B009RC1Q027SBEA",
    "Real_Nonresidential_Structures_Investment": "B009RX1Q020SBEA",
}

for _name, _sid in FRED_SERIES.items():
    _note = ""
    if _sid == "A679RZ2Q224SBEA":
        _note = "Official definition: contribution to percent change in real private fixed investment, not contribution to total GDP growth."
    if _sid == "FEDFUNDS":
        _note = "Candidate U.S. policy-rate series; the final policy-rate definition remains a research-design choice."
    if _sid in {"SP500", "DGS10", "DCOILWTICO", "DTWEXBGS"}:
        _note = "Raw market level/yield retained for H4 candidate coverage; return construction and final asset universe are not imposed here."
    if _sid in {"B009RC1Q027SBEA", "B009RX1Q020SBEA"}:
        _note = "Candidate non-AI investment comparison series for H6; the final exclusion rule for AI/digital investment remains to be specified."
    register(_name, "FRED", _sid, "US", requirement_role="research-document or supporting U.S. series", note=_note)


def fetch_fred_series(series_id: str) -> pd.DataFrame:
    if FRED_API_KEY:
        r = SESSION.get(
            "https://api.stlouisfed.org/fred/series/observations",
            params={
                "series_id": series_id,
                "api_key": FRED_API_KEY,
                "file_type": "json",
            },
            timeout=REQUEST_TIMEOUT,
        )
        r.raise_for_status()
        obs = r.json().get("observations", [])
        df = pd.DataFrame(obs)
        if df.empty:
            return df
        df = df[["date", "value"]].rename(columns={"date": "Date", "value": series_id})
    else:
        # Official FRED graph CSV endpoint does not require an API key.
        r = SESSION.get(
            "https://fred.stlouisfed.org/graph/fredgraph.csv",
            params={"id": series_id},
            timeout=REQUEST_TIMEOUT,
        )
        r.raise_for_status()
        df = pd.read_csv(io.StringIO(r.text))
        if "DATE" in df.columns:
            df = df.rename(columns={"DATE": "Date"})
        elif "observation_date" in df.columns:
            df = df.rename(columns={"observation_date": "Date"})
        elif "date" in df.columns:
            df = df.rename(columns={"date": "Date"})
        else:
            raise ValueError(f"FRED CSV did not contain a recognized date column: {list(df.columns)}")
        if series_id not in df.columns:
            raise ValueError(f"FRED CSV did not contain expected column {series_id}")
        df = df[["Date", series_id]]
    df[series_id] = pd.to_numeric(df[series_id], errors="coerce")
    return df


def collect_fred(writer) -> None:
    print("\n--- FRED ---")
    mode = "API" if FRED_API_KEY else "official keyless CSV fallback"
    for name, series_id in FRED_SERIES.items():
        try:
            df = fetch_fred_series(series_id)
            write_sheet(writer, df, name, f"FRED ({mode})")
        except Exception as exc:
            log(name, "FRED", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


# ---------------------------------------------------------------------------
# World Bank WDI
# ---------------------------------------------------------------------------

WORLDBANK_INDICATORS = {
    "WB_Real_GDP_Level": "NY.GDP.MKTP.KD",
    "WB_Labor_Productivity": "SL.GDP.PCAP.EM.KD",
    "WB_Employment_to_Population": "SL.EMP.TOTL.SP.ZS",
    "WB_Internet_Users_Percent": "IT.NET.USER.ZS",
    "WB_Fixed_Broadband_Per_100": "IT.NET.BBND.P2",
    "WB_R_and_D_Percent_GDP": "GB.XPD.RSDV.GD.ZS",
    "WB_HighTech_Manufactured_Exports": "TX.VAL.TECH.MF.ZS",
    "WB_GDP_Growth_Annual_Percent": "NY.GDP.MKTP.KD.ZG",
    "WB_Unemployment_Percent": "SL.UEM.TOTL.ZS",
    "WB_Inflation_CPI_Percent": "FP.CPI.TOTL.ZG",
    "WB_ICT_Service_Exports_Percent": "BX.GSR.CCIS.ZS",
}

for _name, _sid in WORLDBANK_INDICATORS.items():
    _note = ""
    if _sid == "BX.GSR.CCIS.ZS":
        _note = "Optional separate metric. It is not a substitute for high-technology manufactured export share."
    register(_name, "World Bank WDI", _sid, "Cross-country", "Annual", "research-document / supporting cross-country series", _note)

WORLDBANK_GEM_INDICATORS = {
    "WB_GEM_Industrial_Production_SA": "IPTOTSAKD",
    "WB_GEM_Industrial_Production_NSA": "IPTOTNSKD",
}
for _name, _sid in WORLDBANK_GEM_INDICATORS.items():
    register(
        _name,
        "World Bank Global Economic Monitor (GEM)",
        _sid,
        "Cross-country",
        "Monthly/annual as returned by GEM",
        "industrial-production activity measure",
        "Official GEM series; source id 15 in the World Bank API.",
    )


def fetch_worldbank_entity_metadata() -> tuple[pd.DataFrame, dict[str, dict]]:
    r = SESSION.get(
        "https://api.worldbank.org/v2/country",
        params={"format": "json", "per_page": 400},
        timeout=REQUEST_TIMEOUT,
    )
    r.raise_for_status()
    payload = r.json()
    records = payload[1] if isinstance(payload, list) and len(payload) >= 2 and payload[1] else []
    rows = []
    mapping: dict[str, dict] = {}
    for item in records:
        iso3 = item.get("id", "")
        region = (item.get("region") or {})
        income = (item.get("incomeLevel") or {})
        lending = (item.get("lendingType") or {})
        entity_type = "aggregate" if region.get("id") == "NA" else "country"
        row = {
            "ISO3": iso3,
            "ISO2": item.get("iso2Code", ""),
            "Entity_Name": item.get("name", ""),
            "Entity_Type": entity_type,
            "Region_ID": region.get("id", ""),
            "Region": region.get("value", ""),
            "Income_Level": income.get("value", ""),
            "Lending_Type": lending.get("value", ""),
        }
        rows.append(row)
        mapping[iso3] = row
    return pd.DataFrame(rows), mapping


def fetch_worldbank_indicator(code: str, entity_meta: dict[str, dict]) -> pd.DataFrame:
    page = 1
    rows = []
    while True:
        r = SESSION.get(
            f"https://api.worldbank.org/v2/country/all/indicator/{code}",
            params={"format": "json", "per_page": 20000, "page": page},
            timeout=REQUEST_TIMEOUT,
        )
        r.raise_for_status()
        payload = r.json()
        if not isinstance(payload, list) or len(payload) < 2:
            raise ValueError("Unexpected World Bank response structure")
        meta = payload[0] or {}
        data = payload[1] or []
        rows.extend(data)
        pages = int(meta.get("pages", 1) or 1)
        if page >= pages:
            break
        page += 1
        time.sleep(SLEEP_BETWEEN_CALLS)

    df = pd.json_normalize(rows)
    if df.empty:
        return df
    keep = [
        c for c in ["country.value", "countryiso3code", "date", "value"]
        if c in df.columns
    ]
    df = df[keep].rename(columns={
        "country.value": "Country",
        "countryiso3code": "ISO3",
        "date": "Year",
        "value": code,
    })
    df["Year"] = pd.to_numeric(df["Year"], errors="coerce").astype("Int64")
    df[code] = pd.to_numeric(df[code], errors="coerce")
    df["Entity_Type"] = df["ISO3"].map(lambda x: entity_meta.get(str(x), {}).get("Entity_Type", "unknown"))
    df["Region"] = df["ISO3"].map(lambda x: entity_meta.get(str(x), {}).get("Region", ""))
    df["Income_Level"] = df["ISO3"].map(lambda x: entity_meta.get(str(x), {}).get("Income_Level", ""))
    return df


def longest_contiguous_annual_span(years) -> tuple[int | None, int | None]:
    vals = sorted({int(y) for y in years if pd.notna(y)})
    if not vals:
        return None, None
    best_start = best_end = cur_start = cur_end = vals[0]
    for y in vals[1:]:
        if y == cur_end + 1:
            cur_end = y
        else:
            if (cur_end - cur_start) > (best_end - best_start):
                best_start, best_end = cur_start, cur_end
            cur_start = cur_end = y
    if (cur_end - cur_start) > (best_end - best_start):
        best_start, best_end = cur_start, cur_end
    return best_start, best_end


def add_worldbank_coverage(df: pd.DataFrame, name: str, code: str) -> None:
    if df.empty:
        return
    for iso3, group in df.groupby("ISO3", dropna=False):
        valid = group[numeric_valid(group[code])].copy()
        entity_type = str(group["Entity_Type"].iloc[0]) if "Entity_Type" in group.columns else "unknown"
        if valid.empty:
            _coverage_rows.append({
                "variable_id": code,
                "variable_name": name,
                "country_iso3": iso3,
                "entity_type": entity_type,
                "source": "World Bank WDI",
                "series_id": code,
                "frequency": "annual",
                "first_valid_period": None,
                "last_valid_period": None,
                "valid_observations": 0,
                "expected_periods": None,
                "missing_periods": None,
                "coverage_percent": None,
                "longest_contiguous_start": None,
                "longest_contiguous_end": None,
                "observation_status": "no numeric value returned",
                "comparability_status": "same WDI indicator; no numeric observation for this entity in returned data",
                "availability_status": "missing in returned source data",
                "notes": "",
                "retrieved_at": RETRIEVED_AT,
            })
            continue
        years = pd.to_numeric(valid["Year"], errors="coerce").dropna().astype(int)
        first = int(years.min())
        last = int(years.max())
        expected = last - first + 1
        valid_count = int(years.nunique())
        missing = max(expected - valid_count, 0)
        contiguous_start, contiguous_end = longest_contiguous_annual_span(years.tolist())
        _coverage_rows.append({
            "variable_id": code,
            "variable_name": name,
            "country_iso3": iso3,
            "entity_type": entity_type,
            "source": "World Bank WDI",
            "series_id": code,
            "frequency": "annual",
            "first_valid_period": first,
            "last_valid_period": last,
            "valid_observations": valid_count,
            "expected_periods": expected,
            "missing_periods": missing,
            "coverage_percent": valid_count / expected if expected else None,
            "longest_contiguous_start": contiguous_start,
            "longest_contiguous_end": contiguous_end,
            "observation_status": "observed source values",
            "comparability_status": "same WDI indicator; source/methodology revisions may still require review",
            "availability_status": "available",
            "notes": "Coverage is based on numeric values, not merely first/last rows.",
            "retrieved_at": RETRIEVED_AT,
        })


def collect_worldbank(writer) -> None:
    global _country_name_to_iso3
    print("\n--- World Bank ---")
    entity_df, entity_meta = fetch_worldbank_entity_metadata()
    _country_name_to_iso3 = {
        str(row["Entity_Name"]).strip(): str(row["ISO3"]).strip()
        for _, row in entity_df.iterrows()
        if str(row.get("Entity_Type", "")) == "country"
    }
    write_sheet(writer, entity_df, "WorldBank_Entity_Metadata", "World Bank")
    for name, code in WORLDBANK_INDICATORS.items():
        try:
            df = fetch_worldbank_indicator(code, entity_meta)
            write_sheet(writer, df, name, "World Bank")
            add_worldbank_coverage(df, name, code)
            add_iran_long_rows(df, "World Bank WDI", code, "Year", code, "ISO3", "observed source value")
        except Exception as exc:
            log(name, "World Bank", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


# ---------------------------------------------------------------------------
# World Bank Global Economic Monitor (GEM)
# ---------------------------------------------------------------------------

def _gem_period_key(value):
    text_value = str(value)
    m = re.match(r"^(\d{4})(?:M(\d{1,2})|Q(\d))?$", text_value)
    if not m:
        return (9999, 99, text_value)
    year = int(m.group(1))
    month = int(m.group(2)) if m.group(2) else (int(m.group(3)) * 3 if m.group(3) else 0)
    return (year, month, text_value)


def _normalize_country_label(value) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value).lower()).strip()


def fetch_worldbank_gem_indicator(code: str, entity_meta: dict[str, dict]) -> pd.DataFrame:
    page = 1
    rows = []
    while True:
        r = SESSION.get(
            f"https://api.worldbank.org/v2/country/all/indicator/{code}",
            params={
                "source": 15,
                "format": "json",
                "per_page": 20000,
                "page": page,
            },
            timeout=REQUEST_TIMEOUT,
        )
        r.raise_for_status()
        payload = r.json()
        if not isinstance(payload, list) or len(payload) < 2:
            raise ValueError("Unexpected World Bank GEM response structure")
        meta = payload[0] or {}
        data = payload[1] or []
        rows.extend(data)
        pages = int(meta.get("pages", 1) or 1)
        if page >= pages:
            break
        page += 1
        time.sleep(SLEEP_BETWEEN_CALLS)

    df = pd.json_normalize(rows)
    if df.empty:
        return df
    keep = [
        col for col in ["country.value", "countryiso3code", "date", "value"]
        if col in df.columns
    ]
    df = df[keep].rename(columns={
        "country.value": "Country",
        "countryiso3code": "ISO3",
        "date": "Period",
        "value": code,
    })
    df[code] = pd.to_numeric(df[code], errors="coerce")

    # GEM sometimes returns country names with a blank ISO3 code. Recover ISO3
    # from official World Bank entity metadata so Iran and other countries are
    # auditable instead of silently disappearing from country coverage.
    name_to_iso3 = {
        _normalize_country_label(meta.get("Entity_Name", "")): iso3
        for iso3, meta in entity_meta.items()
        if meta.get("Entity_Type") == "country"
    }
    iso_text = df["ISO3"].astype(str).str.strip()
    missing_iso = iso_text.isin({"", "nan", "None", "<NA>"})
    df.loc[missing_iso, "ISO3"] = df.loc[missing_iso, "Country"].map(
        lambda value: name_to_iso3.get(_normalize_country_label(value), "")
    )
    # Explicit alias for GEM's punctuation-free Iran label.
    iran_alias = df["Country"].astype(str).map(_normalize_country_label) == "iran islamic rep"
    df.loc[iran_alias & (df["ISO3"].astype(str).str.strip() == ""), "ISO3"] = "IRN"

    df["Entity_Type"] = df["ISO3"].map(
        lambda x: entity_meta.get(str(x), {}).get("Entity_Type", "unknown")
    )
    return df


def add_worldbank_gem_coverage(df: pd.DataFrame, name: str, code: str) -> None:
    if df.empty:
        return
    for (iso3, country_name), group in df.groupby(["ISO3", "Country"], dropna=False):
        valid = group[numeric_valid(group[code])].copy()
        entity_type = str(group["Entity_Type"].iloc[0]) if "Entity_Type" in group.columns else "unknown"
        if valid.empty:
            _coverage_rows.append({
                "variable_id": code,
                "variable_name": name,
                "country_iso3": iso3,
                "entity_type": entity_type,
                "source": "World Bank GEM",
                "series_id": code,
                "frequency": "mixed GEM period labels",
                "first_valid_period": None,
                "last_valid_period": None,
                "valid_observations": 0,
                "expected_periods": None,
                "missing_periods": None,
                "coverage_percent": None,
                "longest_contiguous_start": None,
                "longest_contiguous_end": None,
                "observation_status": "entity returned but no numeric industrial-production value",
                "comparability_status": "no usable numeric GEM observation in current extract",
                "availability_status": "missing from current GEM extract",
                "notes": f"World Bank GEM source=15; returned entity name: {country_name}.",
                "retrieved_at": RETRIEVED_AT,
            })
            continue

        periods = sorted(valid["Period"].astype(str).unique().tolist(), key=_gem_period_key)
        _coverage_rows.append({
            "variable_id": code,
            "variable_name": name,
            "country_iso3": iso3,
            "entity_type": entity_type,
            "source": "World Bank GEM",
            "series_id": code,
            "frequency": "mixed GEM period labels",
            "first_valid_period": periods[0],
            "last_valid_period": periods[-1],
            "valid_observations": len(periods),
            "expected_periods": None,
            "missing_periods": None,
            "coverage_percent": None,
            "longest_contiguous_start": None,
            "longest_contiguous_end": None,
            "observation_status": "official source values",
            "comparability_status": "constant-US$ industrial-production series; inspect frequency/base metadata before estimation",
            "availability_status": "available",
            "notes": f"World Bank GEM source=15; returned entity name: {country_name}.",
            "retrieved_at": RETRIEVED_AT,
        })


def collect_worldbank_gem(writer) -> None:
    print("\n--- World Bank GEM ---")
    _, entity_meta = fetch_worldbank_entity_metadata()
    for name, code in WORLDBANK_GEM_INDICATORS.items():
        try:
            df = fetch_worldbank_gem_indicator(code, entity_meta)
            write_sheet(writer, df, name, "World Bank GEM")
            add_worldbank_gem_coverage(df, name, code)
            add_iran_long_rows(
                df,
                "World Bank GEM",
                code,
                "Period",
                code,
                "ISO3",
                "official GEM source value",
            )
        except Exception as exc:
            log(name, "World Bank GEM", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


# ---------------------------------------------------------------------------
# BEA NIPA (optional API key)
# ---------------------------------------------------------------------------

BEA_TABLES = {
    "BEA_RealGDP_PercentChange_T10101": "T10101",
}
register("BEA_RealGDP_PercentChange_T10101", "BEA NIPA", "T10101", "US", "Quarterly", "supporting U.S. GDP data")


def _collect_bea_nipa_public_download(writer) -> bool:
    """Use BEA's public NIPA flat-file downloads when no API key is configured."""
    candidates = [
        ("BEA_NIPA_Section1", "https://apps.bea.gov/national/Release/TXT/Section1All_csv.csv", "csv"),
        ("BEA_NIPA_Section1", "https://apps.bea.gov/national/Release/XLS/Section1All_xls.xlsx", "xlsx"),
        ("BEA_NIPA_Annual_Bulk", "https://apps.bea.gov/national/Release/TXT/NipaDataA.zip", "zip"),
        ("BEA_NIPA_Quarterly_Bulk", "https://apps.bea.gov/national/Release/TXT/NipaDataQ.zip", "zip"),
    ]
    errors = []
    for label, url, kind in candidates:
        try:
            r = SESSION.get(url, timeout=REQUEST_TIMEOUT)
            r.raise_for_status()
            if not r.content:
                raise ValueError("empty response")
            if kind == "csv":
                try:
                    df = pd.read_csv(io.BytesIO(r.content), low_memory=False)
                except Exception:
                    # BEA section flat files include title/metadata rows whose
                    # field counts differ from the tabular rows. csv.reader
                    # preserves the official rows without dropping that context.
                    text_data = r.content.decode("utf-8-sig", errors="replace")
                    rows = list(csv.reader(io.StringIO(text_data)))
                    if not rows:
                        raise ValueError("CSV response contained no rows")
                    width = max(len(row) for row in rows)
                    padded = [row + [""] * (width - len(row)) for row in rows]
                    df = pd.DataFrame(
                        padded,
                        columns=[f"Field_{i+1}" for i in range(width)],
                    )
                write_sheet(writer, df, label, "BEA NIPA public download")
                log("BEA NIPA public-download route", "BEA", "OK", url)
                return True
            if kind == "xlsx":
                sheets = pd.read_excel(io.BytesIO(r.content), sheet_name=None)
                total = 0
                for sheet_name, df in sheets.items():
                    if df is None or df.empty:
                        continue
                    write_sheet(writer, df, f"{label}_{sheet_name}", "BEA NIPA public download")
                    total += len(df)
                if total:
                    log("BEA NIPA public-download route", "BEA", "OK", f"{url} ({total} rows)")
                    return True
                raise ValueError("workbook contained no rows")
            if kind == "zip":
                with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
                    members = [
                        m for m in zf.namelist()
                        if m.lower().endswith((".csv", ".txt"))
                        and not m.endswith("/")
                    ]
                    if not members:
                        raise ValueError("zip contained no CSV/TXT files")
                    total = 0
                    for member in members[:12]:
                        with zf.open(member) as fh:
                            raw = fh.read()
                        try:
                            df = pd.read_csv(io.BytesIO(raw), low_memory=False)
                        except Exception:
                            df = pd.read_csv(io.BytesIO(raw), sep="\t", low_memory=False)
                        if df.empty:
                            continue
                        base = member.rsplit("/", 1)[-1].rsplit(".", 1)[0]
                        write_sheet(writer, df, f"{label}_{base}", "BEA NIPA public download")
                        total += len(df)
                    if total:
                        log("BEA NIPA public-download route", "BEA", "OK", f"{url} ({total} rows)")
                        return True
                    raise ValueError("zip files were readable but contained no rows")
        except Exception as exc:
            errors.append(f"{url}: {exc}")
    log(
        "All BEA NIPA tables",
        "BEA",
        "ERROR",
        "Public download attempts failed: " + " | ".join(errors),
    )
    return False


def collect_bea(writer) -> None:
    print("\n--- BEA NIPA ---")
    if not BEA_API_KEY:
        _collect_bea_nipa_public_download(writer)
        return
    for name, table in BEA_TABLES.items():
        try:
            r = SESSION.get(
                "https://apps.bea.gov/api/data",
                params={
                    "UserID": BEA_API_KEY,
                    "method": "GetData",
                    "DatasetName": "NIPA",
                    "TableName": table,
                    "Frequency": "Q",
                    "Year": "ALL",
                    "ResultFormat": "JSON",
                },
                timeout=REQUEST_TIMEOUT,
            )
            r.raise_for_status()
            result = r.json().get("BEAAPI", {}).get("Results", {})
            if "Error" in result:
                raise ValueError(str(result["Error"]))
            data = result.get("Data", [])
            write_sheet(writer, pd.DataFrame(data), name, "BEA")
        except Exception as exc:
            log(name, "BEA", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


# ---------------------------------------------------------------------------
# BLS - chunked history
# ---------------------------------------------------------------------------

BLS_SERIES = {
    "BLS_Unemployment_Rate": "LNS14000000",
    "BLS_Total_Nonfarm_Payrolls": "CES0000000001",
    "BLS_CPI_All_Items": "CUUR0000SA0",
    "BLS_Computer_Systems_Design_Emp": "CES6054150001",
    "BLS_PPI_Information_Sector": "PCUAINFO-AINFO-",
    "BLS_PPI_Data_Processing_Hosting": "PCU518210518210",
    "BLS_PPI_Wired_Telecom": "PCU517311517311",
    "BLS_PPI_Wireless_Telecom": "PCU517312517312",
    "BLS_PPI_Total_Manufacturing": "PCUOMFGOMFG",
    "BLS_PPI_Total_Mining": "PCUOMINOMIN",
    "BLS_PPI_Transportation_Warehousing": "PCUATRNWRATRNWR",
    "BLS_PPI_Wholesale_Trade": "PCUAWHLTRAWHLTR",
    "BLS_PPI_Retail_Trade": "PCUARETTRARETTR",
}

# FRED republishes these BLS series through its official keyless CSV endpoint.
# Using the mirror in keyless CI avoids the BLS anonymous daily-request ceiling
# while preserving the BLS series provenance.
BLS_FRED_MIRRORS = {
    "BLS_Unemployment_Rate": "UNRATE",
    "BLS_Total_Nonfarm_Payrolls": "PAYEMS",
    "BLS_CPI_All_Items": "CPIAUCNS",
    "BLS_Computer_Systems_Design_Emp": "CES6054150001",
    "BLS_PPI_Information_Sector": "PCUAINFOAINFO",
    "BLS_PPI_Data_Processing_Hosting": "PCU518210518210",
    "BLS_PPI_Wired_Telecom": "PCU517311517311",
    "BLS_PPI_Wireless_Telecom": "PCU517312517312",
    "BLS_PPI_Total_Manufacturing": "PCUOMFGOMFG",
    "BLS_PPI_Total_Mining": "PCUOMINOMIN",
    "BLS_PPI_Transportation_Warehousing": "PCUATRNWRATRNWR",
    "BLS_PPI_Wholesale_Trade": "PCUAWHLTRAWHLTR",
    "BLS_PPI_Retail_Trade": "PCUARETTRARETTR",
}

for _name, _sid in BLS_SERIES.items():
    mirror = BLS_FRED_MIRRORS.get(_name, "")
    register(
        _name,
        "BLS (official series; FRED mirror used when keyless)",
        _sid,
        "US",
        "Monthly",
        "research-document / supporting labor-price series",
        f"Keyless FRED mirror series: {mirror}" if mirror else "",
    )

# 1913 is the earliest start among the currently configured BLS series set (CPI).
# It is a source-specific floor, not a project-wide research start year.
BLS_HISTORY_START = 1913


def _collect_bls_via_fred_mirror(writer) -> None:
    for name, bls_sid in BLS_SERIES.items():
        fred_sid = BLS_FRED_MIRRORS.get(name)
        if not fred_sid:
            log(name, "BLS/FRED mirror", "ERROR", "No FRED mirror mapping configured")
            continue
        try:
            raw = fetch_fred_series(fred_sid)
            if raw.empty:
                write_sheet(writer, raw, name, "FRED mirror of BLS")
                continue
            out = raw.rename(columns={fred_sid: "value_numeric"}).copy()
            out["BLS_Series_ID"] = bls_sid
            out["FRED_Mirror_ID"] = fred_sid
            out["Source_Note"] = "Official BLS-origin series retrieved via Federal Reserve FRED mirror"
            write_sheet(writer, out, name, "FRED mirror of BLS")
        except Exception as exc:
            log(name, "BLS/FRED mirror", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


def collect_bls(writer) -> None:
    print("\n--- BLS ---")

    # GitHub Actions normally has no BLS registration key. The public BLS API
    # has a low anonymous daily request ceiling, so repeated professional
    # builds can regress. Use FRED's official mirror in that case.
    if not BLS_API_KEY:
        _collect_bls_via_fred_mirror(writer)
        return

    span = 20
    by_series: dict[str, list[dict]] = {sid: [] for sid in BLS_SERIES.values()}

    chunk_start = BLS_HISTORY_START
    while chunk_start <= END_YEAR:
        chunk_end = min(chunk_start + span - 1, END_YEAR)
        body = {
            "seriesid": list(BLS_SERIES.values()),
            "startyear": str(chunk_start),
            "endyear": str(chunk_end),
            "registrationKey": BLS_API_KEY,
        }

        try:
            r = SESSION.post(
                "https://api.bls.gov/publicAPI/v2/timeseries/data/",
                json=body,
                headers={"Content-type": "application/json"},
                timeout=REQUEST_TIMEOUT,
            )
            r.raise_for_status()
            payload = r.json()
            if payload.get("status") != "REQUEST_SUCCEEDED":
                raise ValueError(str(payload.get("message", "")))
            for series in payload.get("Results", {}).get("series", []):
                sid = series.get("seriesID", "")
                if sid in by_series:
                    by_series[sid].extend(series.get("data", []))
        except Exception as exc:
            log(f"BLS chunk {chunk_start}-{chunk_end}", "BLS", "ERROR", str(exc))
        chunk_start = chunk_end + 1
        time.sleep(SLEEP_BETWEEN_CALLS)

    id_to_name = {v: k for k, v in BLS_SERIES.items()}
    for sid, rows in by_series.items():
        name = id_to_name[sid]
        df = pd.DataFrame(rows)
        if not df.empty:
            if "year" in df.columns and "period" in df.columns:
                df = df.drop_duplicates(subset=["year", "period"], keep="first")
                df = df.sort_values(["year", "period"], ascending=[False, False])
            if "value" in df.columns:
                df["value_numeric"] = pd.to_numeric(df["value"], errors="coerce")
        write_sheet(writer, df, name, "BLS")


# ---------------------------------------------------------------------------
# OECD supplementary CLI
# ---------------------------------------------------------------------------

register(
    "OECD_Composite_Leading_Indicator",
    "OECD",
    "OECD.SDD.STES,DSD_STES@DF_CLI",
    "OECD/selected areas",
    "Monthly",
    "supplementary business-cycle indicator; NOT a substitute for Going Digital/OECD.AI metrics",
)


def collect_oecd(writer) -> None:
    print("\n--- OECD ---")
    try:
        url = (
            "https://sdmx.oecd.org/public/rest/data/"
            "OECD.SDD.STES,DSD_STES@DF_CLI/.M.LI...AA...H"
        )
        r = SESSION.get(
            url,
            params={
                "dimensionAtObservation": "AllDimensions",
                "format": "csvfilewithlabels",
            },
            timeout=REQUEST_TIMEOUT,
        )
        r.raise_for_status()
        df = pd.read_csv(io.StringIO(r.text))
        write_sheet(writer, df, "OECD_Composite_Leading_Indicator", "OECD")
    except Exception as exc:
        log("OECD_Composite_Leading_Indicator", "OECD", "ERROR", str(exc))


# ---------------------------------------------------------------------------
# IMF DataMapper
# ---------------------------------------------------------------------------

IMF_INDICATORS = {
    "IMF_AI_Preparedness_Index": "AI_PI",
    "IMF_Real_GDP_Growth": "NGDP_RPCH",
    "IMF_Inflation_Percent": "PCPIPCH",
}
register("IMF_AI_Preparedness_Index", "IMF DataMapper / AIPI", "AI_PI", "Cross-country", "2023 snapshot", "AI preparedness; not AI investment/adoption")
register("IMF_Real_GDP_Growth", "IMF DataMapper / WEO", "NGDP_RPCH", "Cross-country", "Annual", "macro control/outcome")
register("IMF_Inflation_Percent", "IMF DataMapper / WEO", "PCPIPCH", "Cross-country", "Annual", "macro price control")


def add_imf_coverage(df: pd.DataFrame, name: str, code: str) -> None:
    if df.empty:
        return
    for country, group in df.groupby("Country_Code", dropna=False):
        valid = group[numeric_valid(group[code])]
        if valid.empty:
            continue
        years = pd.to_numeric(valid["Year"], errors="coerce").dropna().astype(int)
        returned_first = int(years.min())
        returned_last = int(years.max())
        # DataMapper output here does not carry a reliable actual/estimate/forecast flag.
        # To avoid calling future values observed history, leave last_valid_period blank.
        _coverage_rows.append({
            "variable_id": code,
            "variable_name": name,
            "country_iso3": country,
            "entity_type": "country_or_imf_entity",
            "source": "IMF DataMapper",
            "series_id": code,
            "frequency": "annual_or_snapshot",
            "first_valid_period": returned_first,
            "last_valid_period": None,
            "valid_observations": int(len(valid)),
            "expected_periods": None,
            "missing_periods": None,
            "coverage_percent": None,
            "longest_contiguous_start": None,
            "longest_contiguous_end": None,
            "observation_status": "returned periods; actual/estimate/forecast status not retained",
            "comparability_status": "IMF WEO returned periods include future values; observed/estimate/forecast split unresolved",
            "availability_status": "available but observed endpoint unresolved",
            "notes": f"Returned through {returned_last}; do not treat returned endpoint as last observed year.",
            "retrieved_at": RETRIEVED_AT,
        })


def collect_imf(writer) -> None:
    print("\n--- IMF ---")
    for name, code in IMF_INDICATORS.items():
        try:
            r = SESSION.get(
                f"https://www.imf.org/external/datamapper/api/v1/{code}",
                timeout=REQUEST_TIMEOUT,
            )
            r.raise_for_status()
            values = r.json().get("values", {}).get(code, {})
            records = []
            for country, year_map in values.items():
                for year, value in year_map.items():
                    try:
                        y = int(year)
                    except Exception:
                        y = year
                    period_flag = (
                        "future_year_returned"
                        if isinstance(y, int) and y > END_YEAR
                        else "current_or_historical_year_returned"
                    )
                    records.append({
                        "Country_Code": country,
                        "Year": y,
                        code: value,
                        "Period_Flag": period_flag,
                    })
            df = pd.DataFrame(records)
            if not df.empty:
                df[code] = pd.to_numeric(df[code], errors="coerce")
            write_sheet(writer, df, name, "IMF")
            add_imf_coverage(df, name, code)
            add_iran_long_rows(
                df,
                "IMF DataMapper",
                code,
                "Year",
                code,
                "Country_Code",
                "returned; actual/estimate/forecast classification not retained",
            )
        except Exception as exc:
            log(name, "IMF", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


# ---------------------------------------------------------------------------
# BEA Digital Economy latest official workbook
# ---------------------------------------------------------------------------

BEA_DIGITAL_ECONOMY_URL = (
    "https://bea.gov/sites/default/files/2023-12/"
    "DigitalEconomy_2017-2022.xlsx"
)
register(
    "BEA_Digital_Economy_Satellite_Account",
    "BEA Digital Economy",
    "2017-2022 official workbook",
    "US",
    "Annual",
    "digital economy share/output/value added; account discontinued after 2023 release",
)


def collect_bea_digital_economy(writer) -> None:
    print("\n--- BEA Digital Economy ---")
    try:
        r = SESSION.get(BEA_DIGITAL_ECONOMY_URL, timeout=REQUEST_TIMEOUT)
        r.raise_for_status()
        sheets = pd.read_excel(io.BytesIO(r.content), sheet_name=None)
        wrote = 0
        for idx, (sheet_name, df) in enumerate(sheets.items(), start=1):
            if df is None or df.empty:
                continue
            out_name = f"BEA_DE_{idx:02d}_{sheet_name}"
            write_sheet(writer, df, out_name, "BEA Digital Economy 2017-2022")
            wrote += 1
        if wrote == 0:
            log("BEA_Digital_Economy", "BEA Digital Economy", "EMPTY", "Workbook contained no non-empty sheets")
    except Exception as exc:
        log("BEA_Digital_Economy", "BEA Digital Economy", "ERROR", str(exc))


# ---------------------------------------------------------------------------
# Epoch AI model-compute data
# ---------------------------------------------------------------------------

EPOCH_DATASETS = {
    "Epoch_Frontier_AI_Models": "https://epoch.ai/data/frontier_ai_models.csv",
    "Epoch_Notable_AI_Models": "https://epoch.ai/data/notable_ai_models.csv",
}
register(
    "Epoch_Frontier_AI_Models",
    "Epoch AI",
    "frontier_ai_models.csv",
    "Model-level / organization / country fields as published",
    "Irregular event-level",
    "frontier-model training compute and related model characteristics",
    "Epoch defines frontier models as models in the top five of training compute at release.",
)
register(
    "Epoch_Notable_AI_Models",
    "Epoch AI",
    "notable_ai_models.csv",
    "Model-level / organization / country fields as published",
    "Irregular event-level",
    "supporting AI model history",
)


def collect_epoch_ai(writer) -> None:
    print("\n--- Epoch AI ---")
    for name, url in EPOCH_DATASETS.items():
        try:
            r = SESSION.get(url, timeout=REQUEST_TIMEOUT)
            r.raise_for_status()
            df = pd.read_csv(io.StringIO(r.text))
            write_sheet(writer, df, name, "Epoch AI")
        except Exception as exc:
            log(name, "Epoch AI", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


# ---------------------------------------------------------------------------
# AIOE / AIIE occupational and industry exposure data
# ---------------------------------------------------------------------------

AIOE_FILES = {
    "AIOE_Base": "https://raw.githubusercontent.com/AIOE-Data/AIOE/main/AIOE_DataAppendix.xlsx",
    "AIOE_GenAI_Language": "https://raw.githubusercontent.com/AIOE-Data/AIOE/main/Language%20Modeling%20AIOE%20and%20AIIE.xlsx",
    "AIOE_GenAI_Image": "https://raw.githubusercontent.com/AIOE-Data/AIOE/main/Image%20Generation%20AIOE%20and%20AIIE.xlsx",
}
register(
    "AIOE_AIIE_AIGE",
    "Felten-Raj-Seamans / AIOE-Data",
    "AIOE_DataAppendix.xlsx",
    "US occupation / industry / county",
    "Cross-sectional exposure index",
    "H2 occupational, industry, and geographic AI exposure",
    "Base workbook contains AIOE by SOC, AIIE by NAICS, and AIGE by FIPS.",
)
register(
    "AIOE_Generative_AI_Extensions",
    "AIOE-Data",
    "Language Modeling / Image Generation AIOE and AIIE",
    "US occupation / industry",
    "Cross-sectional exposure index",
    "H2 generative-AI exposure extensions",
)


def collect_aioe(writer) -> None:
    print("\n--- AIOE / AIIE ---")
    for file_label, url in AIOE_FILES.items():
        try:
            r = SESSION.get(url, timeout=REQUEST_TIMEOUT)
            r.raise_for_status()
            book = pd.read_excel(io.BytesIO(r.content), sheet_name=None)
            wrote = 0
            for idx, (sheet_name, df) in enumerate(book.items(), start=1):
                if df is None or df.empty:
                    continue
                write_sheet(
                    writer,
                    df,
                    f"{file_label}_{idx:02d}_{sheet_name}",
                    "AIOE-Data",
                )
                wrote += 1
            if wrote == 0:
                log(file_label, "AIOE-Data", "EMPTY", "No non-empty sheets found")
        except Exception as exc:
            log(file_label, "AIOE-Data", "ERROR", str(exc))
        time.sleep(SLEEP_BETWEEN_CALLS)


# ---------------------------------------------------------------------------
# CSET / ETO Country AI Activity Metrics
# ---------------------------------------------------------------------------

CSET_ETO_FILES = {
    "CSET_AI_Patent_Applications": "patents_yearly_applications.csv",
    "CSET_AI_Investment_Disclosed": "companies_yearly_disclosed.csv",
    "CSET_AI_Investment_Estimated": "companies_yearly_estimated.csv",
}
CSET_ETO_RECORD_API = "https://zenodo.org/api/records/22772306"
register(
    "CSET_AI_Patent_Applications",
    "CSET / ETO Country AI Activity Metrics",
    "patents_yearly_applications.csv",
    "Cross-country",
    "Annual",
    "AI patent filings",
    "Recent patent years can be materially incomplete; retain the source complete flag.",
)
register(
    "CSET_AI_Investment_Disclosed",
    "CSET / ETO Country AI Activity Metrics",
    "companies_yearly_disclosed.csv",
    "Cross-country",
    "Annual",
    "private-market AI investment",
    "Millions USD; equity investment into privately held AI-related companies.",
)
register(
    "CSET_AI_Investment_Estimated",
    "CSET / ETO Country AI Activity Metrics",
    "companies_yearly_estimated.csv",
    "Cross-country",
    "Annual",
    "estimated private-market AI investment",
    "Millions USD; preserve the source complete flag and distinguish estimated from disclosed values.",
)


def _zenodo_file_url(file_meta: dict) -> str:
    links = file_meta.get("links") or {}
    return links.get("content") or links.get("self") or links.get("download") or ""


def add_cset_country_coverage(
    df: pd.DataFrame,
    variable_name: str,
    series_id: str,
    value_col: str,
) -> None:
    if df.empty or "country" not in df.columns or "year" not in df.columns:
        return
    work = df.copy()
    work["year"] = pd.to_numeric(work["year"], errors="coerce")
    work[value_col] = pd.to_numeric(work[value_col], errors="coerce")
    for country, group in work.groupby("country", dropna=False):
        valid = group[group[value_col].notna() & group["year"].notna()].copy()
        if valid.empty:
            continue
        years = sorted(valid["year"].astype(int).unique().tolist())
        first = years[0]
        last = years[-1]
        expected = last - first + 1
        valid_years = len(years)
        missing = max(expected - valid_years, 0)
        cont_start, cont_end = longest_contiguous_annual_span(years)
        complete_rows = None
        if "complete" in valid.columns:
            complete_rows = int(valid["complete"].fillna(False).astype(bool).sum())
        note = (
            f"Field-disaggregated source. Coverage counts unique years with any returned row; "
            f"do not sum fields without reviewing overlap. Complete-flagged rows: {complete_rows}."
        )
        source_country = str(country).strip()
        mapped_iso3 = _country_name_to_iso3.get(source_country)
        if source_country == "Iran":
            mapped_iso3 = "IRN"
        _coverage_rows.append({
            "variable_id": series_id,
            "variable_name": variable_name,
            "country_iso3": mapped_iso3 or source_country,
            "entity_type": "country" if mapped_iso3 else "country_or_group_as_named_by_CSET",
            "source": "CSET / ETO Country AI Activity Metrics",
            "series_id": series_id,
            "frequency": "annual",
            "first_valid_period": first,
            "last_valid_period": last,
            "valid_observations": valid_years,
            "expected_periods": expected,
            "missing_periods": missing,
            "coverage_percent": valid_years / expected if expected else None,
            "longest_contiguous_start": cont_start,
            "longest_contiguous_end": cont_end,
            "observation_status": "source rows; retain complete flag for material completeness",
            "comparability_status": "field-disaggregated CSET/ETO definitions; recent patent years may be incomplete",
            "availability_status": "available",
            "notes": note,
            "retrieved_at": RETRIEVED_AT,
        })

    iran = work[work["country"].astype(str) == "Iran"].copy()
    for _, row in iran.iterrows():
        _iran_rows.append({
            "source": "CSET / ETO Country AI Activity Metrics",
            "series_id": series_id,
            "country_iso3": "IRN",
            "period": row.get("year"),
            "value": row.get(value_col),
            "field": row.get("field"),
            "observation_status": (
                "source row; complete=" + str(row.get("complete"))
            ),
            "retrieved_at": RETRIEVED_AT,
        })


def collect_cset_eto(writer) -> None:
    print("\n--- CSET / ETO Country AI Activity ---")
    try:
        meta_r = SESSION.get(CSET_ETO_RECORD_API, timeout=REQUEST_TIMEOUT)
        meta_r.raise_for_status()
        meta = meta_r.json()
        files = meta.get("files") or []
        manifest = pd.DataFrame([
            {
                "key": f.get("key"),
                "size": f.get("size"),
                "checksum": f.get("checksum"),
                "download_url": _zenodo_file_url(f),
            }
            for f in files
        ])
        if not manifest.empty:
            write_sheet(writer, manifest, "CSET_ETO_File_Manifest", "CSET / ETO Zenodo")
        else:
            log("CSET_ETO_File_Manifest", "CSET / ETO", "EMPTY", "Zenodo record returned no files")

        targets_by_filename = {filename: name for name, filename in CSET_ETO_FILES.items()}
        found: set[str] = set()

        for f in files:
            key = str(f.get("key") or "")
            url = _zenodo_file_url(f)
            if not url:
                continue

            base = key.rsplit("/", 1)[-1]
            if base in targets_by_filename:
                name = targets_by_filename[base]
                try:
                    r = SESSION.get(url, timeout=REQUEST_TIMEOUT)
                    r.raise_for_status()
                    df = pd.read_csv(io.BytesIO(r.content))
                    write_sheet(writer, df, name, "CSET / ETO")
                    value_col = {
                        "patents_yearly_applications.csv": "num_patent_applications",
                        "companies_yearly_disclosed.csv": "disclosed_investment",
                        "companies_yearly_estimated.csv": "estimated_investment",
                    }[base]
                    add_cset_country_coverage(df, name, base, value_col)
                    found.add(base)
                except Exception as exc:
                    log(name, "CSET / ETO", "ERROR", str(exc))
                continue

            if key.lower().endswith(".zip"):
                try:
                    r = SESSION.get(url, timeout=REQUEST_TIMEOUT)
                    r.raise_for_status()
                    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
                        members = zf.namelist()
                        for target_file, name in targets_by_filename.items():
                            match = next(
                                (m for m in members if m.rsplit("/", 1)[-1] == target_file),
                                None,
                            )
                            if match is None:
                                continue
                            with zf.open(match) as fh:
                                df = pd.read_csv(fh)
                            write_sheet(writer, df, name, "CSET / ETO")
                            value_col = {
                                "patents_yearly_applications.csv": "num_patent_applications",
                                "companies_yearly_disclosed.csv": "disclosed_investment",
                                "companies_yearly_estimated.csv": "estimated_investment",
                            }[target_file]
                            add_cset_country_coverage(df, name, target_file, value_col)
                            found.add(target_file)
                except Exception as exc:
                    log(f"CSET archive {key}", "CSET / ETO", "ERROR", str(exc))

        available_keys = ", ".join(str(f.get("key") or "") for f in files)
        for filename, name in targets_by_filename.items():
            if filename not in found:
                log(
                    name,
                    "CSET / ETO",
                    "EMPTY",
                    f"Target file not found in current Zenodo record. Available keys: {available_keys}",
                )
    except Exception as exc:
        log("CSET / ETO record metadata", "CSET / ETO", "ERROR", str(exc))



# ---------------------------------------------------------------------------
# World Bank Digital Adoption Index and Census BTOS
# ---------------------------------------------------------------------------

def collect_worldbank_dai(writer) -> None:
    """Download the official long-form Digital Adoption Index workbook."""
    page_url = "https://www.worldbank.org/en/publication/wdr2016/Digital-Adoption-Index"
    try:
        page = SESSION.get(page_url, timeout=REQUEST_TIMEOUT)
        page.raise_for_status()
        hrefs = re.findall(r'href="([^"]+\\.xlsx(?:\\?[^"]*)?)"', page.text, flags=re.I)
        if not hrefs:
            hrefs = re.findall(r"href='([^']+\\.xlsx(?:\\?[^']*)?)'", page.text, flags=re.I)
        if hrefs:
            file_url = urljoin(page_url, hrefs[0])
        else:
            # The current World Bank page renders the download control without
            # exposing the target href in the static HTML. This is the World Bank
            # file referenced for the same DAI long-form download.
            file_url = "https://thedocs.worldbank.org/en/doc/625521534508595697-0050022018/original/DAIforweb.xlsx"
        r = SESSION.get(file_url, timeout=REQUEST_TIMEOUT)
        r.raise_for_status()
        sheets = pd.read_excel(io.BytesIO(r.content), sheet_name=None)
        wrote = 0
        for sheet_name, df in sheets.items():
            if df is None or df.empty:
                continue
            write_sheet(writer, df, f"DAI_{sheet_name}", "World Bank DAI")
            wrote += len(df)

            # The official DAI workbook is a sparse 2014/2016 cross-section.
            # Preserve Iran explicitly without fabricating annual observations.
            lower_cols = {str(col).strip().lower(): col for col in df.columns}
            country_col = lower_cols.get("country")
            year_col = lower_cols.get("year")
            dai_col = lower_cols.get("digital adoption index")
            if country_col and year_col and dai_col:
                iran_mask = df[country_col].astype(str).str.strip().str.startswith("Iran")
                iran_part = df[iran_mask].copy()
                valid_iran = iran_part[pd.to_numeric(iran_part[dai_col], errors="coerce").notna()].copy()
                for _, row in valid_iran.iterrows():
                    _iran_rows.append({
                        "source": "World Bank DAI",
                        "series_id": "DAI",
                        "country_iso3": "IRN",
                        "period": row.get(year_col),
                        "value": row.get(dai_col),
                        "observation_status": "official sparse cross-section",
                        "retrieved_at": RETRIEVED_AT,
                    })
                if not valid_iran.empty:
                    years = sorted(pd.to_numeric(valid_iran[year_col], errors="coerce").dropna().astype(int).unique().tolist())
                    _coverage_rows.append({
                        "variable_id": "DAI",
                        "variable_name": "World Bank Digital Adoption Index",
                        "country_iso3": "IRN",
                        "entity_type": "country",
                        "source": "World Bank DAI",
                        "series_id": "DAI",
                        "frequency": "sparse cross-section",
                        "first_valid_period": min(years),
                        "last_valid_period": max(years),
                        "valid_observations": len(years),
                        "expected_periods": 2,
                        "missing_periods": max(2 - len(years), 0),
                        "coverage_percent": len(years) / 2,
                        "longest_contiguous_start": None,
                        "longest_contiguous_end": None,
                        "observation_status": "official 2014/2016 observations",
                        "comparability_status": "official DAI values; not an annual panel",
                        "availability_status": "available",
                        "notes": "Do not interpolate the missing calendar year between official DAI waves.",
                        "retrieved_at": RETRIEVED_AT,
                    })
        if wrote:
            log("Digital Adoption Index workbook", "World Bank DAI", "OK", f"{wrote} rows across {len(sheets)} source sheets", wrote)
        else:
            log("Digital Adoption Index workbook", "World Bank DAI", "EMPTY")
    except Exception as exc:
        log("Digital Adoption Index workbook", "World Bank DAI", "ERROR", str(exc))


def collect_btos(writer) -> None:
    """Download current official BTOS workbooks that contain national and AI-question data."""
    files = [
        ("BTOS_National", "https://www.census.gov/hfp/btos/downloads/National.xlsx"),
        ("BTOS_AI_Core_Questions", "https://www.census.gov/hfp/btos/downloads/AI%20Core%20Questions.xlsx"),
    ]
    for label, url in files:
        try:
            r = SESSION.get(url, timeout=REQUEST_TIMEOUT)
            r.raise_for_status()
            sheets = pd.read_excel(io.BytesIO(r.content), sheet_name=None)
            total = 0
            for sheet_name, df in sheets.items():
                if df is None or df.empty:
                    continue
                write_sheet(writer, df, f"{label}_{sheet_name}", "U.S. Census BTOS")
                total += len(df)
            if total:
                log(label, "U.S. Census BTOS", "OK", f"{total} rows across {len(sheets)} source sheets", total)
            else:
                log(label, "U.S. Census BTOS", "EMPTY")
        except Exception as exc:
            log(label, "U.S. Census BTOS", "ERROR", str(exc))


# ---------------------------------------------------------------------------
# Manual / not-yet-automated sources
# ---------------------------------------------------------------------------

MANUAL_SOURCES = [
    {
        "Variable": "Census BTOS – AI use by businesses",
        "URL": "https://www.census.gov/programs-surveys/btos.html",
        "Status": "official national and AI-question workbooks collected",
        "Note": "Raw BTOS workbooks are included; preserve question/version, sector, geography and sample metadata in any later modeling.",
    },
    {
        "Variable": "World Bank – Digital Adoption Index",
        "URL": "https://www.worldbank.org/en/publication/wdr2016/Digital-Adoption-Index",
        "Status": "official 2014/2016 workbook collected",
        "Note": "Sparse cross-section (2014 and 2016), not an annual panel; do not interpolate into fake yearly observations.",
    },
    {
        "Variable": "Stanford HAI – AI Index economy data",
        "URL": "https://hai.stanford.edu/ai-index/2026-ai-index-report",
        "Status": "data acquisition still required",
        "Note": "Select exact public tables and preserve the underlying provider/methodology.",
    },
    {
        "Variable": "Epoch AI – model aggregation into country-time panels",
        "URL": "https://epoch.ai/data/ai-models",
        "Status": "raw Frontier/Notable model files automated; aggregation methodology still required",
        "Note": "Keep model-level observations distinct from any later country/year aggregation.",
    },
    {
        "Variable": "CSET / ETO Country Activity Tracker",
        "URL": "https://cat.eto.tech/",
        "Status": "AI patent and private-investment country files collected",
        "Note": "Disclosed and estimated investment files are kept separate; preserve completeness flags and source lags.",
    },
    {
        "Variable": "AIOE / AIIE / AIGE employment matching",
        "URL": "https://github.com/AIOE-Data/AIOE",
        "Status": "raw exposure workbooks automated; compatible employment histories/crosswalk still required",
        "Note": "Do not treat exposure scores alone as H2-ready without matched SOC/NAICS employment/output data.",
    },
    {
        "Variable": "OECD Going Digital / OECD.AI",
        "URL": "https://goingdigital.oecd.org/indicators",
        "Status": "OECD digital-intensity sector taxonomy documented; additional OECD.AI indicator acquisition still optional",
        "Note": "The ISIC Rev.4 intensity mapping is embedded separately; CLI remains only a business-cycle supplement.",
    },
    {
        "Variable": "Iran industrial production",
        "URL": "https://data.imf.org/Datasets/PI",
        "Status": "IMF Iran annual/quarterly series identifiers documented; GEM Iran entity returned but current extract has no numeric values",
        "Note": "The gap remains open until numeric Iran IPI values are obtained through a reproducible official route; do not substitute manufacturing value added.",
    },
    {
        "Variable": "Iran monetary-policy rate",
        "URL": "https://www.imf.org/-/media/files/publications/cr/2017/cr1763.pdf",
        "Status": "CBI repo/corridor observations documented as operational candidates",
        "Note": "IMF documents that Iran historically lacked a conventional single policy rate; use repo/interbank/corridor measures only after the model definition is fixed.",
    },
    {
        "Variable": "Iran digital economy share of GDP",
        "URL": "https://nezamat.ir/%D9%86%D8%B8%D8%A7%D9%85-%D9%86%D8%A7%D9%85%D9%87-%D8%B3%D9%86%D8%AC%D8%B4-%D9%88-%D8%A8%D9%87-%D8%B1%D9%88%D8%B2%D8%B1%D8%B3%D8%A7%D9%86%DB%8C-%D8%B3%D9%87%D9%85-%D8%B2%DB%8C%D8%B3%D8%AA-%D8%A8%D9%88/",
        "Status": "two officially reported observations documented; underlying SCI table still preferred if released",
        "Note": "Keep the reported values separate from interpolated annual panels; no interpolation is authorized.",
    },
]


OECD_DIGITAL_INTENSITY_TAXONOMY = [
    ["01-03", "Agriculture, forestry and fishing", "Low"],
    ["05-09", "Mining and quarrying", "Low"],
    ["10-12", "Food products, beverages and tobacco", "Low"],
    ["13-15", "Textiles, wearing apparel, leather", "Medium-low"],
    ["16-18", "Wood and paper products; printing", "Medium-high"],
    ["19-23", "Chemicals, rubber, plastics, fuel and other non-metallic mineral products", "Medium-low"],
    ["24-25", "Basic metals and fabricated metal products", "Medium-low"],
    ["26-28", "Machinery and equipment", "Medium-high"],
    ["29-30", "Transport equipment", "High"],
    ["31-33", "Furniture, other manufacturing, repair/installation", "Medium-high"],
    ["35-39", "Electricity, gas, water, waste", "Low"],
    ["41-43", "Construction", "Low"],
    ["45-47", "Wholesale/retail trade and motor-vehicle repair", "Medium-high"],
    ["49-53", "Transportation and storage", "Low"],
    ["55-56", "Accommodation and food service", "Low"],
    ["58-60", "Publishing, audiovisual and broadcasting", "Medium-high"],
    ["61", "Telecommunications", "High"],
    ["62-63", "IT and other information services", "High"],
    ["64-66", "Financial and insurance activities", "High"],
    ["68", "Real estate activities", "Low"],
    ["69-82", "Professional, scientific, technical, administrative and support services", "High"],
    ["84", "Public administration and defence", "Medium-high"],
    ["85", "Education", "Medium-low"],
    ["86-88", "Human health and social work", "Medium-low"],
    ["90-93", "Arts, entertainment and recreation", "Medium-high"],
    ["94-96", "Other service activities", "High"],
]

IRAN_REPORTED_DIGITAL_ECONOMY = [
    {
        "Iranian_Year": 1400,
        "Approx_Gregorian_Year": 2021,
        "Digital_Economy_Share_GDP_Percent": 4.49,
        "Evidence_Type": "official estimate reported by deputy ICT minister; underlying SCI table not directly retrieved",
        "Measurement_Authority_Reported": "Statistical Center of Iran + Ministry of ICT",
        "Source_URL": "https://ecomotive.ir/2026/06/17/%D8%A7%D9%82%D8%AA%D8%B5%D8%A7%D8%AF-%D8%AF%DB%8C%D8%AC%DB%8C%D8%AA%D8%A7%D9%84-%DB%B7-%DB%B1%DB%B4-%D8%AF%D8%B1%D8%B5%D8%AF-%D8%A7%D8%B2-%D8%A7%D9%82%D8%AA%D8%B5%D8%A7%D8%AF-%DA%A9%D9%84-%DA%A9%D8%B4/",
        "Method_Framework_URL": "https://nezamat.ir/%D9%86%D8%B8%D8%A7%D9%85-%D9%86%D8%A7%D9%85%D9%87-%D8%B3%D9%86%D8%AC%D8%B4-%D9%88-%D8%A8%D9%87-%D8%B1%D9%88%D8%B2%D8%B1%D8%B3%D8%A7%D9%86%DB%8C-%D8%B3%D9%87%D9%85-%D8%B2%DB%8C%D8%B3%D8%AA-%D8%A8%D9%88/",
    },
    {
        "Iranian_Year": 1403,
        "Approx_Gregorian_Year": 2024,
        "Digital_Economy_Share_GDP_Percent": 7.14,
        "Evidence_Type": "official estimate reported by deputy ICT minister; underlying SCI table not directly retrieved",
        "Measurement_Authority_Reported": "Statistical Center of Iran + Ministry of ICT",
        "Source_URL": "https://ecomotive.ir/2026/06/17/%D8%A7%D9%82%D8%AA%D8%B5%D8%A7%D8%AF-%D8%AF%DB%8C%D8%AC%DB%8C%D8%AA%D8%A7%D9%84-%DB%B7-%DB%B1%DB%B4-%D8%AF%D8%B1%D8%B5%D8%AF-%D8%A7%D8%B2-%D8%A7%D9%82%D8%AA%D8%B5%D8%A7%D8%AF-%DA%A9%D9%84-%DA%A9%D8%B4/",
        "Method_Framework_URL": "https://nezamat.ir/%D9%86%D8%B8%D8%A7%D9%85-%D9%86%D8%A7%D9%85%D9%87-%D8%B3%D9%86%D8%AC%D8%B4-%D9%88-%D8%A8%D9%87-%D8%B1%D9%88%D8%B2%D8%B1%D8%B3%D8%A7%D9%86%DB%8C-%D8%B3%D9%87%D9%85-%D8%B2%DB%8C%D8%B3%D8%AA-%D8%A8%D9%88/",
    },
]

IRAN_INDUSTRIAL_PRODUCTION_SOURCE_MAP = [
    {
        "Source": "IMF Production Indexes (PI)",
        "Series_ID": "IRN.IND.IX.A",
        "Frequency": "Annual",
        "Reported_Start": 1948,
        "Direct_Values_In_Workbook": "NO",
        "Status": "series identifier located in an IMF-derived database mirror; direct IMF values still need a reproducible official download route",
        "Primary_URL": "https://data.imf.org/Datasets/PI",
        "Identifier_Evidence_URL": "https://net-imf.tedc.org.tw/",
    },
    {
        "Source": "IMF Production Indexes (PI)",
        "Series_ID": "IRN.IND.IX.Q",
        "Frequency": "Quarterly",
        "Reported_Start": 1948,
        "Direct_Values_In_Workbook": "NO",
        "Status": "series identifier located in an IMF-derived database mirror; direct IMF values still need a reproducible official download route",
        "Primary_URL": "https://data.imf.org/Datasets/PI",
        "Identifier_Evidence_URL": "https://net-imf.tedc.org.tw/",
    },
    {
        "Source": "World Bank Global Economic Monitor",
        "Series_ID": "IPTOTSAKD",
        "Frequency": "as returned by GEM",
        "Reported_Start": None,
        "Direct_Values_In_Workbook": "CHECK COVERAGE",
        "Status": "Iran entity is returned by GEM; run 22 showed no numeric Iran observations, so this is not treated as filled",
        "Primary_URL": "https://api.worldbank.org/v2/country/all/indicator/IPTOTSAKD?source=15&format=json",
        "Identifier_Evidence_URL": "",
    },
    {
        "Source": "Iran Parliament Research Center industrial-production index",
        "Series_ID": "definition/source route only",
        "Frequency": "monthly/periodic source dependent",
        "Reported_Start": None,
        "Direct_Values_In_Workbook": "NO",
        "Status": "World Bank Iran Economic Monitor documents an IPI covering more than 302 listed industrial companies; raw series not yet automated",
        "Primary_URL": "https://www.worldbank.org/en/country/iran/publication/economic-monitor",
        "Identifier_Evidence_URL": "",
    },
]

IRAN_MONETARY_POLICY_CANDIDATES = [
    {
        "Iranian_Date": "1402-11 onward (documented in 1403 report)",
        "Metric": "Minimum repo rate in weekly open-market operations",
        "Percent": 23.0,
        "Interpretation": "Parliament Research Center documents 23% as the effective weekly-auction target; retain as an operational proxy, not a universal conventional policy rate.",
        "Source_URL": "https://doi.org/10.22034/report.mrc.2024.1403.32.7.20582",
    },
    {
        "Iranian_Date": "1402-11 onward (documented in 1403 report)",
        "Metric": "Interest-rate corridor floor",
        "Percent": 17.0,
        "Interpretation": "Lower bound of the CBI interest-rate corridor documented by the Parliament Research Center.",
        "Source_URL": "https://doi.org/10.22034/report.mrc.2024.1403.32.7.20582",
    },
    {
        "Iranian_Date": "1402-11 onward (documented in 1403 report)",
        "Metric": "Interest-rate corridor ceiling",
        "Percent": 24.0,
        "Interpretation": "Upper bound of the CBI interest-rate corridor documented by the Parliament Research Center.",
        "Source_URL": "https://doi.org/10.22034/report.mrc.2024.1403.32.7.20582",
    },
    {
        "Iranian_Date": "1403-10-10",
        "Metric": "Minimum repo rate in open-market operation",
        "Percent": 23.0,
        "Interpretation": "Operational OMO rate candidate; do not label as a unique policy rate without professor/model definition.",
        "Source_URL": "https://42946232.khabarban.com/",
    },
    {
        "Iranian_Date": "1403-10-10",
        "Metric": "Standing lending facility / corridor ceiling",
        "Percent": 24.0,
        "Interpretation": "Ceiling of the CBI interest-rate corridor reported with the OMO release.",
        "Source_URL": "https://42946232.khabarban.com/",
    },
    {
        "Iranian_Date": "1404-06-31",
        "Metric": "Minimum repo rate in open-market operation",
        "Percent": 23.0,
        "Interpretation": "Operational OMO rate candidate; not automatically equivalent to a conventional single policy rate.",
        "Source_URL": "https://wikibanki.com/%DA%AF%D8%B2%D8%A7%D8%B1%D8%B4-%D8%B9%D9%85%D9%84%DB%8C%D8%A7%D8%AA-%D8%B9%D8%B1%D8%B6%D9%87-%D8%A7%D8%AC%D8%B1%D8%A7%DB%8C%DB%8C-%D8%B3%DB%8C%D8%A7%D8%B3%D8%AA-%D9%BE%D9%88%D9%84%DB%8C-%D8%A8%D8%A7%D9%86%DA%A9-%D9%85%D8%B1/",
    },
]

RESEARCH_REQUIREMENTS = [
    ["BC01","Business cycle","Real GDP level","US high-frequency + country equivalents","FRED GDPC1; WDI NY.GDP.MKTP.KD","US + Iran + cross-country WDI level collected","available"],
    ["BC02","Business cycle","Real GDP growth","Core outcome/control","WDI NY.GDP.MKTP.KD.ZG; IMF NGDP_RPCH; FRED A191RL1Q225SBEA","Iran + cross-country + US collected","available"],
    ["BC03","Business cycle","Industrial production index","Monthly activity control","FRED INDPRO; World Bank GEM IPTOTSAKD/IPTOTNSKD; IMF Production Indexes","US collected; GEM Iran entity returned but no numeric Iran values in run 22; IMF Iran annual/quarterly identifiers documented for follow-up","partial; Iran numeric series still missing"],
    ["BC04","Business cycle","Unemployment rate","Labor control","WDI SL.UEM.TOTL.ZS; BLS/FRED","Iran annual + cross-country + US monthly collected","available"],
    ["BC05","Business cycle","Nonfarm payroll employment","US labor control","FRED PAYEMS / BLS CES0000000001","US collected; not an Iran series","available US-only"],
    ["BC06","Business cycle","Recession indicator/dating","Regime variable","FRED USREC / NBER","US collected; not an Iran series","available US-only"],
    ["BC07","Business cycle","Policy interest rate","Monetary-policy interaction","FRED FEDFUNDS candidate; CBI open-market-operation corridor; IMF monetary framework notes","US candidate collected; Iran repo/corridor snapshots documented, but no single conventional policy-rate series is imposed","partial / proxy available"],
    ["DE01","Digital economy","Internet penetration","Digital intensity","WDI IT.NET.USER.ZS","Iran + cross-country collected","available"],
    ["DE02","Digital economy","Fixed broadband subscriptions per 100","Digital intensity","WDI IT.NET.BBND.P2","Iran + cross-country collected","available"],
    ["DE03","Digital economy","High-technology exports (% manufactured exports)","Digital trade/intensity","WDI TX.VAL.TECH.MF.ZS","Iran + cross-country collected","available"],
    ["DE04","Digital economy","R&D expenditure (% GDP)","Innovation intensity","WDI GB.XPD.RSDV.GD.ZS","Iran + cross-country collected","available"],
    ["DE05","Digital economy","Digital economy share of GDP","National-account digital share","BEA Digital Economy + national equivalents","US BEA 2017-2022 collected; Iran officially reported estimates documented for 1400 and 1403 with methodology caveat","partial; sparse Iran observations"],
    ["DE06","Digital economy","Composite digital-transformation index","Cross-sectional/panel digital intensity","World Bank DAI; OECD Going Digital; DESI","World Bank DAI 2014/2016 workbook collected, including Iran; OECD Going Digital remains a supplementary route","partial; DAI is sparse cross-section"],
    ["AI01","AI-specific","Private investment in AI","AI capital intensity","CSET/ETO; Stanford HAI; OECD.AI","CSET/ETO disclosed + estimated annual country data collected with completeness flags","available from CSET/ETO; other source variants optional"],
    ["AI02","AI-specific","AI patent filings","Innovation proxy","CSET/ETO; OECD.AI","CSET/ETO annual country/field patent applications collected; completeness flags retained","available from CSET/ETO"],
    ["AI03","AI-specific","Frontier-model training compute","Physical AI-capital proxy","Epoch AI","Frontier + notable model raw data collected","available model-level"],
    ["AI04","AI-specific","AI venture capital/private-market investment","AI investment proxy","CSET/ETO; Stanford HAI; OECD.AI","CSET/ETO private-market disclosed/estimated investment collected","available from CSET/ETO"],
    ["AI05","AI-specific","Information-processing equipment & software investment","Long-run digital-capex proxy","FRED A679RC1Q027SBEA","US collected","available US-only"],
    ["AI06","AI-specific","Business AI use/adoption","High-frequency adoption","US Census BTOS","Official national and AI-question workbooks collected","available US-only"],
    ["LB01","Labor/structure","Information/high-tech industry employment","Structural labor measure","FRED USINFO; BLS CES6054150001","US collected","available US-only"],
    ["LB02","Labor/structure","Occupational/industry/geographic AI exposure","H2 exposure variable","AIOE/AIIE/AIGE","base + generative-language + image exposure workbooks collected","available exposure data"],
    ["LB03","Labor/structure","Labor productivity","H1 outcome","FRED OPHNFB; WDI SL.GDP.PCAP.EM.KD","US + Iran + cross-country annual productivity collected","available"],
    ["LB04","Labor/structure","Total factor productivity","H1 outcome","FRED MFPPBS","US collected","available US-only"],
    ["LB05","Labor/structure","Unit labor costs","Cost/price mechanism","FRED ULCNFB","US collected","available US-only"],
    ["LB06","Labor/structure","Job openings","Matching efficiency","FRED JTSJOL / JTS540099JOL","US collected","available US-only"],
    ["LB07","Labor/structure","Labor turnover","Matching efficiency","FRED JTSHIL / JTSQUL / JTSTSL","US collected","available US-only"],
    ["PR01","Prices","Headline CPI/inflation","Price control/outcome","WDI FP.CPI.TOTL.ZG; FRED CPIAUCSL; BLS CUUR0000SA0","Iran + cross-country + US collected","available"],
    ["PR02","Prices","Core CPI / core PCE","Underlying inflation","FRED CPILFESL / PCEPILFE","US collected","available US-only"],
    ["PR03","Prices","Sectoral CPI/PPI by industry","H3 sectoral persistence","BLS series via official FRED mirror + OECD digital-intensity taxonomy","ICT PPI plus manufacturing, mining, transport/warehousing, wholesale and retail PPI collected reproducibly; OECD ISIC Rev.4 digital-intensity mapping embedded","partial; NAICS-to-ISIC crosswalk and persistence definition remain"],
    ["PR04","Prices","Inflation volatility","Derived rolling standard deviation","Derived from price series","raw inputs partly collected; window/frequency not approved","derived methodology pending"],
    ["FN01","Finance/uncertainty","Economic Policy Uncertainty","H6 uncertainty measure","FRED USEPUINDXM / USEPUINDXD","US monthly + daily collected","available US-only"],
    ["FN02","Finance/uncertainty","VIX","Financial interaction","FRED VIXCLS","US market collected","available US-only"],
    ["FN03","Finance/uncertainty","Financial conditions","Robustness control","FRED NFCI","US collected","available US-only"],
    ["FN04","Finance/uncertainty","Multiple asset returns","H4 cross-asset correlation","FRED market candidates","S&P 500, 10-year Treasury yield, WTI oil and broad dollar series collected as raw candidates","partial; return construction and final asset universe pending"],
    ["INV01","Investment","Information-processing investment contribution","H5","FRED A679RZ2Q224SBEA","US collected; official definition is contribution to real private fixed-investment growth, NOT total GDP growth","available with definition correction"],
    ["H6A","Events","Major AI release calendar","H6 event alignment","Epoch AI model release dates","Model-level release dates collected; rule for what counts as a major capability event remains to be fixed","partial"],
    ["H6B","Investment","Non-AI investment","H6 comparison outcome","FRED/BEA investment components","Nonresidential structures series collected as a comparison candidate; exact exclusion rule remains undefined","partial / definition pending"],
]

HYPOTHESIS_DATA_MATRIX = [
    ["H1","Digital/AI-capex growth vs productivity with lags","Digital capex; labor productivity; TFP; long history","US capex/productivity/TFP + WDI country labor productivity + CSET country AI investment collected","partial","Country-year overlap and the final digital/AI-capex definition must be audited; no estimation authorized"],
    ["H2","Employment volatility/output volatility by AI exposure","AIOE/AIIE; occupation/industry employment; output; crosswalk","AIOE/AIIE exposure + WDI employment-to-population/output series (including Iran) collected","partial","Compatible occupation/industry historical employment-output crosswalk is still needed"],
    ["H3","Sectoral price persistence vs digital intensity","Sectoral CPI/PPI; sector digital intensity; persistence definition","Monthly BLS-origin PPI series retrieved via FRED mirror across ICT plus several broader sectors; OECD ISIC Rev.4 intensity taxonomy documented","partial","NAICS-to-ISIC crosswalk and persistence-window definition remain; no estimator is imposed"],
    ["H4","Cross-asset correlation/volatility during AI-capex growth","Multiple asset returns; AI/digital capex; window definition","VIX, capex and several market-level candidate series collected","partial","Return construction, final asset universe and frequency are still to be specified"],
    ["H5","Post-2023 volatility of information-processing investment contribution","Contribution series; comparator components; long history","A679RZ2Q224SBEA collected","partial","Research document mislabels this as total-GDP contribution; comparator components must be selected"],
    ["H6","EPU around major AI releases followed by weaker non-AI investment","EPU; release calendar; non-AI investment","EPU, Epoch model release dates and nonresidential-investment candidates collected","partial","Major-event selection and the final non-AI investment definition are still to be specified"],
]


def write_requirement_matrices(writer) -> None:
    req_cols = ["ID","Family","Variable","Research_Role","Candidate_Source","Current_Evidence","Availability_Status"]
    hyp_cols = ["Hypothesis","Question","Required_Data","Current_Evidence","Data_Readiness","Remaining_Gap"]
    excel_safe_dataframe(pd.DataFrame(RESEARCH_REQUIREMENTS, columns=req_cols)).to_excel(
        writer, sheet_name="Requirements Matrix", index=False
    )
    excel_safe_dataframe(pd.DataFrame(HYPOTHESIS_DATA_MATRIX, columns=hyp_cols)).to_excel(
        writer, sheet_name="Hypothesis Matrix", index=False
    )
    excel_safe_dataframe(
        pd.DataFrame(
            OECD_DIGITAL_INTENSITY_TAXONOMY,
            columns=["ISIC_Rev4_Divisions", "Sector", "Digital_Intensity"],
        )
    ).to_excel(writer, sheet_name="OECD Digital Intensity", index=False)
    excel_safe_dataframe(pd.DataFrame(IRAN_REPORTED_DIGITAL_ECONOMY)).to_excel(
        writer, sheet_name="Iran Digital Economy", index=False
    )
    excel_safe_dataframe(pd.DataFrame(IRAN_INDUSTRIAL_PRODUCTION_SOURCE_MAP)).to_excel(
        writer, sheet_name="Iran IPI Source Map", index=False
    )
    excel_safe_dataframe(pd.DataFrame(IRAN_MONETARY_POLICY_CANDIDATES)).to_excel(
        writer, sheet_name="Iran Monetary Rates", index=False
    )



def format_client_workbook(writer) -> None:
    """Apply restrained, conventional formatting for a client-facing research workbook."""
    wb = writer.book
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.properties.title = "Digital Economy and Business Cycle Data"
    wb.properties.subject = "Source data, coverage and variable notes"

    # A short front sheet helps a reviewer understand the file without reading the script.
    if "Read Me" in wb.sheetnames:
        del wb["Read Me"]
    ws = wb.create_sheet("Read Me", 0)
    ws["A1"] = "Digital Economy, Artificial Intelligence and Business Cycle Data"
    ws["A3"] = "Version"
    ws["B3"] = "28 September 2026"
    ws["A4"] = "Scope"
    ws["B4"] = "Public-source data collection and coverage review. Iran is retained explicitly wherever the source reports it."
    ws["A5"] = "Data handling"
    ws["B5"] = "Values remain in source units. Missing values are kept blank. No interpolation or synthetic observations are added."
    ws["A6"] = "Forecasts"
    ws["B6"] = "Returned future periods are kept separate from certified observed coverage where the source does not expose observation-status metadata."
    ws["A7"] = "Research stage"
    ws["B7"] = "This file is a data delivery and coverage workbook. It does not contain econometric estimates."
    ws["A9"] = "Useful sheets"
    ws["B9"] = "Iran Data; Iran Coverage; Iran Digital Economy; Iran IPI Source Map; Iran Monetary Rates; Country Summary; Coverage; Variable Index; Research Questions; OECD Digital Intensity; Data Dictionary; Source Follow-up; Source Log; Method Notes"
    ws["A11"] = "Source note"
    ws["B11"] = "Source URLs and series identifiers are retained in the data dictionary and source-specific sheets."
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 105
    ws["A1"].font = Font(bold=True, size=16, color="1F1F1F")
    for r in range(3, 12):
        ws[f"A{r}"].font = Font(bold=True, color="1F4E78")
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.sheet_view.showGridLines = False

    # Rename only the client-facing control sheets. Source-data sheets keep their series names.
    rename_map = {
        "Requirements Matrix": "Variable Index",
        "Hypothesis Matrix": "Research Questions",
        "Manual Sources": "Source Follow-up",
        "Variable Dictionary": "Data Dictionary",
        "Collection Status": "Source Log",
        "Audit Notes": "Method Notes",
    }
    for old, new in rename_map.items():
        if old in wb.sheetnames and new not in wb.sheetnames:
            wb[old].title = new

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_border = Border(bottom=Side(style="thin", color="A6A6A6"))

    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = False
        if sh.max_row < 1 or sh.max_column < 1:
            continue
        if sh.title != "Read Me":
            sh.freeze_panes = "A2"
            sh.auto_filter.ref = sh.dimensions
            for cell in sh[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = header_border
            sh.row_dimensions[1].height = 24

            # Widths are based on the header and a small sample, capped to keep large raw sheets readable.
            sample_end = min(sh.max_row, 80)
            for col_idx in range(1, sh.max_column + 1):
                letter = get_column_letter(col_idx)
                max_len = 0
                for row_idx in range(1, sample_end + 1):
                    value = sh.cell(row=row_idx, column=col_idx).value
                    if value is None:
                        continue
                    max_len = max(max_len, len(str(value)))
                sh.column_dimensions[letter].width = min(max(max_len + 2, 11), 38)

    # Put the review sheets directly after Read Me; leave all source sheets behind them.
    preferred = ["Read Me", "Iran Data", "Iran Coverage", "Iran Digital Economy", "Iran IPI Source Map", "Iran Monetary Rates",
                 "Country Summary", "Coverage", "Variable Index", "Research Questions", "OECD Digital Intensity",
                 "Data Dictionary", "Method Notes", "Source Follow-up", "Source Log"]
    ordered = []
    seen = set()
    for name in preferred:
        if name in wb.sheetnames:
            ordered.append(wb[name])
            seen.add(name)
    ordered.extend([sh for sh in wb.worksheets if sh.title not in seen])
    wb._sheets = ordered


def main() -> None:
    print("=" * 78)
    print("Starting revised data collection and coverage audit")
    print("=" * 78)

    try:
        with pd.ExcelWriter(TEMP_OUTPUT_FILE, engine="openpyxl") as writer:
            collect_fred(writer)
            collect_worldbank(writer)
            collect_worldbank_gem(writer)
            collect_bea(writer)
            collect_bls(writer)
            collect_oecd(writer)
            collect_imf(writer)
            collect_bea_digital_economy(writer)
            collect_epoch_ai(writer)
            collect_aioe(writer)
            collect_cset_eto(writer)
            collect_worldbank_dai(writer)
            collect_btos(writer)
            write_requirement_matrices(writer)

            excel_safe_dataframe(pd.DataFrame(MANUAL_SOURCES)).to_excel(
                writer, sheet_name="Manual Sources", index=False
            )

            coverage_df = pd.DataFrame(_coverage_rows)
            if not coverage_df.empty:
                excel_safe_dataframe(coverage_df).to_excel(writer, sheet_name="Coverage", index=False)
                iran_cov = coverage_df[
                    coverage_df["country_iso3"].astype(str) == "IRN"
                ].copy()
                excel_safe_dataframe(iran_cov).to_excel(writer, sheet_name="Iran Coverage", index=False)

                # Client-facing country summary: Iran is mandatory, while all other
                # countries are kept whenever official source coverage exists.
                country_cov = coverage_df[
                    coverage_df["entity_type"].astype(str) == "country"
                ].copy()
                country_cov["valid_observations_num"] = pd.to_numeric(
                    country_cov["valid_observations"], errors="coerce"
                ).fillna(0)
                country_cov = country_cov[country_cov["valid_observations_num"] > 0]

                if not country_cov.empty:
                    flag_ids = {
                        "GDP_Growth": "NY.GDP.MKTP.KD.ZG",
                        "Internet": "IT.NET.USER.ZS",
                        "Broadband": "IT.NET.BBND.P2",
                        "R_and_D": "GB.XPD.RSDV.GD.ZS",
                        "HighTech_Exports": "TX.VAL.TECH.MF.ZS",
                        "Labor_Productivity": "SL.GDP.PCAP.EM.KD",
                        "Employment_Ratio": "SL.EMP.TOTL.SP.ZS",
                        "Industrial_Production_GEM": "IPTOTSAKD",
                        "AI_Patents": "patents_yearly_applications.csv",
                    }
                    investment_ids = {
                        "companies_yearly_disclosed.csv",
                        "companies_yearly_estimated.csv",
                    }
                    summary_rows = []
                    for iso3, group in country_cov.groupby("country_iso3", dropna=False):
                        ids = set(group["variable_id"].astype(str))
                        first_vals = pd.to_numeric(group["first_valid_period"], errors="coerce").dropna()
                        last_vals = pd.to_numeric(group["last_valid_period"], errors="coerce").dropna()
                        row = {
                            "country_iso3": iso3,
                            "mandatory_iran": "YES" if str(iso3) == "IRN" else "",
                            "variables_with_data": int(group["variable_id"].nunique()),
                            "sources_with_data": int(group["source"].nunique()),
                            "first_any_period": int(first_vals.min()) if not first_vals.empty else None,
                            "last_any_period": int(last_vals.max()) if not last_vals.empty else None,
                        }
                        for label, variable_id in flag_ids.items():
                            row[f"has_{label}"] = "YES" if variable_id in ids else ""
                        row["has_AI_Investment"] = "YES" if ids.intersection(investment_ids) else ""
                        summary_rows.append(row)

                    country_summary = pd.DataFrame(summary_rows)
                    country_summary["_iran_sort"] = (
                        country_summary["country_iso3"].astype(str) != "IRN"
                    ).astype(int)
                    country_summary = country_summary.sort_values(
                        ["_iran_sort", "variables_with_data", "country_iso3"],
                        ascending=[True, False, True],
                    ).drop(columns=["_iran_sort"])
                    excel_safe_dataframe(country_summary).to_excel(
                        writer, sheet_name="Country Summary", index=False
                    )

            iran_df = pd.DataFrame(_iran_rows)
            if not iran_df.empty:
                excel_safe_dataframe(iran_df).to_excel(writer, sheet_name="Iran Data", index=False)

            excel_safe_dataframe(pd.DataFrame(_dictionary_rows).drop_duplicates()).to_excel(
                writer, sheet_name="Variable Dictionary", index=False
            )
            excel_safe_dataframe(pd.DataFrame(_log_rows)).to_excel(
                writer, sheet_name="Collection Status", index=False
            )

            method_notes = pd.DataFrame([
                {
                    "Topic": "Iran coverage",
                    "Note": "Iran is shown separately when source data are available. U.S.-only series remain labeled as U.S. data.",
                },
                {
                    "Topic": "Historical range",
                    "Note": "Each source is collected over its available history where practical; a single start year is not forced across all series.",
                },
                {
                    "Topic": "Future periods",
                    "Note": "Future periods returned by IMF are retained but are not counted as observed history unless the source status can be verified.",
                },
                {
                    "Topic": "Missing values",
                    "Note": "Blank observations remain blank. No interpolation or zero filling is used to make the panel appear complete.",
                },
                {
                    "Topic": "Estimation",
                    "Note": "The current workbook is limited to data collection, source notes and coverage. Econometric estimation is outside this delivery stage.",
                },
            ])
            excel_safe_dataframe(method_notes).to_excel(writer, sheet_name="Audit Notes", index=False)
            format_client_workbook(writer)

        os.replace(TEMP_OUTPUT_FILE, OUTPUT_FILE)
    finally:
        if os.path.exists(TEMP_OUTPUT_FILE):
            try:
                os.remove(TEMP_OUTPUT_FILE)
            except OSError:
                pass

    print("=" * 78)
    print(f"Finished: {OUTPUT_FILE}")
    print("=" * 78)


if __name__ == "__main__":
    main()
